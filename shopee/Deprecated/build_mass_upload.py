import argparse
from pathlib import Path
from typing import Dict, List, Optional, Set, cast

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def iter_rows_as_dict(ws: Worksheet, header_row: int, start_row: int) -> List[Dict[str, str]]:
    headers: List[str] = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        headers.append("" if v is None else str(v).strip())

    rows: List[Dict[str, str]] = []
    for r in range(start_row, ws.max_row + 1):
        empty = True
        item: Dict[str, str] = {}
        for c, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(r, c).value
            s = "" if v is None else str(v).strip()
            if s:
                empty = False
            item[h] = s
        if empty:
            continue
        rows.append(item)
    return rows


def build_col_index_from_header(ws: Worksheet, header_row: int) -> Dict[str, int]:
    idx: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is None:
            continue
        name = str(v).strip()
        if name and name not in idx:
            idx[name] = c
    return idx


def get_writable_cell(ws: Worksheet, row: int, col: int) -> Cell:
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
                return cast(Cell, ws.cell(mr.min_row, mr.min_col))
        return cast(Cell, cell)
    return cast(Cell, cell)


def compose_description(item: Dict[str, str]) -> str:
    # 依你的 test.xlsx 欄位拼一個基本描述：特色標語 + 商品特色
    parts: List[str] = []
    slogan = item.get("商品特色標語", "").strip()
    feats = item.get("\n商品特色", "").strip()
    if slogan:
        parts.append(slogan)
    if feats:
        parts.append(feats)
    return "\n".join([p for p in parts if p]).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input-xlsx", type=Path, required=True, help="來源 test.xlsx")
    ap.add_argument("--template", type=Path, required=True, help="Shopee 大量上架範本（可 .xlsm 或 .xlsx）")
    ap.add_argument("--output-xlsx", type=Path, required=True, help="輸出 .xlsx（後台上傳用）")
    ap.add_argument("--input-sheet", type=str, default="", help="來源工作表名稱（留空=第一張）")
    ap.add_argument("--template-sheet", type=str, default="Template", help="範本工作表（預設=Template）")
    ap.add_argument("--input-header-row", type=int, default=1, help="來源表頭列（預設=1）")
    ap.add_argument("--skip-row-2", action="store_true", help="略過第二列（欄位說明）")
    ap.add_argument("--template-header-row", type=int, default=1, help="範本表頭列（預設=1，欄位代碼列）")
    ap.add_argument("--template-data-start-row", type=int, default=7, help="範本資料起始列（預設=7）")
    args = ap.parse_args()

    # 1) 讀來源
    src_wb: Workbook = openpyxl.load_workbook(args.input_xlsx, data_only=True)
    src_ws: Worksheet = src_wb[args.input_sheet] if args.input_sheet else src_wb.worksheets[0]

    input_start_row = 3 if args.skip_row_2 else (args.input_header_row + 1)
    src_rows = iter_rows_as_dict(src_ws, header_row=args.input_header_row, start_row=input_start_row)
    if not src_rows:
        raise ValueError("source xlsx has no data rows (after skipping).")

    # 2) 讀範本（最後要輸出 xlsx，所以不 keep_vba）
    tpl_wb: Workbook = openpyxl.load_workbook(args.template)
    if args.template_sheet not in tpl_wb.sheetnames:
        raise ValueError(
            f"template sheet not found: {args.template_sheet}. "
            f"Available: {tpl_wb.sheetnames}"
        )
    tpl_ws: Worksheet = tpl_wb[args.template_sheet]

    tpl_col_idx = build_col_index_from_header(tpl_ws, args.template_header_row)

    # 3) 最小可用欄位對照：test.xlsx -> template field code
    #    你之後可改成讀 mapping.json（但先讓它能產生「有內容」的 xlsx）
    mapping: Dict[str, str] = {
        "銷售品名(相同當作同一賣場)": "ps_product_name",
        "ERP品號": "ps_sku_parent_short",
        "國際條碼": "ps_product_code",  # 若範本沒有此欄位會自動略過
        "商品材積\n長(cm)": "ps_length",
        "商品材積\n寬(cm)": "ps_width",
        "商品材積\n高(cm)": "ps_height",
        "重量(KG)": "ps_weight",
        "價格\n售價": "ps_price",
        # 規格（1、2）
        "規格名稱1": "et_title_variation_1",
        "規格內容1": "et_title_option_for_variation_1",
        "規格名稱2": "et_title_variation_2",
        "規格內容2": "et_title_option_for_variation_2",
    }

    # 4) 清空資料列（只清我們要寫到的欄位；從 data_start_row 開始）
    target_cols: Set[int] = set()
    for tpl_field in mapping.values():
        c = tpl_col_idx.get(tpl_field)
        if c:
            target_cols.add(c)

    # 額外我們會寫入描述/類別
    for extra in ("ps_product_description", "ps_category"):
        c = tpl_col_idx.get(extra)
        if c:
            target_cols.add(c)

    if target_cols:
        for r in range(args.template_data_start_row, tpl_ws.max_row + 1):
            for c in sorted(target_cols):
                cell = get_writable_cell(tpl_ws, r, c)
                cell.value = None

    # 5) 寫入
    write_row = args.template_data_start_row
    for item in src_rows:
        # 類別：你目前還沒做 category mapping，先留空或先用固定值（之後接 category_list.xls）
        # item["平台"] 之類目前不影響蝦皮範本
        if "ps_category" in tpl_col_idx:
            get_writable_cell(tpl_ws, write_row, tpl_col_idx["ps_category"]).value = ""

        # 商品描述：用特色拼
        if "ps_product_description" in tpl_col_idx:
            desc = compose_description(item)
            get_writable_cell(tpl_ws, write_row, tpl_col_idx["ps_product_description"]).value = desc

        # 其他欄位：照 mapping 填
        for src_h, tpl_field in mapping.items():
            v = str(item.get(src_h, "")).strip()
            if not v:
                continue
            col = tpl_col_idx.get(tpl_field)
            if not col:
                continue
            get_writable_cell(tpl_ws, write_row, col).value = v

        write_row += 1

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    tpl_wb.save(args.output_xlsx)
    print(f"[OK] wrote: {args.output_xlsx}")


if __name__ == "__main__":
    main()
