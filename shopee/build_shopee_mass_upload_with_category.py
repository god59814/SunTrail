import argparse
import math
import re
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any

import openpyxl
import pandas as pd
import gspread
from google.oauth2.service_account import Credentials
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from llm_category_service import predict_category, preload_index


# -----------------------------
# Source helpers
# -----------------------------
SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]


def _normalize_header(h: Any) -> str:
    s = "" if h is None else str(h)
    return s.strip().replace("\n", " ")


def _dedupe_headers(raw_headers: List[Any]) -> List[str]:
    used: Dict[str, int] = {}
    headers: List[str] = []
    for i, h in enumerate(raw_headers, start=1):
        base = _normalize_header(h) or f"__col{i}"
        if base not in used:
            used[base] = 1
            headers.append(base)
        else:
            used[base] += 1
            headers.append(f"{base}__{used[base]}")
    return headers


def read_sheet_records(
    credentials_path: str,
    spreadsheet_id: str,
    worksheet_name: str,
    header_row: int,
    data_start_row: int,
) -> List[Dict[str, str]]:
    creds = Credentials.from_service_account_file(credentials_path, scopes=SCOPES)
    client = gspread.authorize(creds)
    ws = client.open_by_key(spreadsheet_id).worksheet(worksheet_name)
    all_values = ws.get_all_values()
    if len(all_values) < header_row:
        return []

    raw_headers = all_values[header_row - 1]
    headers = _dedupe_headers(raw_headers)
    out: List[Dict[str, str]] = []
    for r in range(data_start_row - 1, len(all_values)):
        row = all_values[r]
        if len(row) < len(headers):
            row = row + [""] * (len(headers) - len(row))
        else:
            row = row[: len(headers)]
        if not any(str(x).strip() for x in row):
            continue
        out.append({k: ("" if v is None else str(v).strip()) for k, v in zip(headers, row)})
    return out


def load_gsheet_pm_source_config() -> Dict[str, Any]:
    project_root = Path(__file__).resolve().parents[1]
    cfg_path = project_root / "done" / "shared" / "config" / "gsheet_pm_source.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    cred_raw = str(cfg.get("cred_path", "") or "").strip()
    cred_path = Path(cred_raw)
    if not cred_path.is_absolute():
        cred_path = (project_root / "done" / cred_path).resolve()
    cfg["cred_abs_path"] = str(cred_path)
    return cfg


def gsheet_row_to_shopee_row(src: Dict[str, str]) -> Dict[str, str]:
    return {
        "銷售品名(相同當作同一賣場)": str(src.get("sale_product_name", "") or ""),
        "商品特色標語": str(src.get("slogan", "") or ""),
        "\n商品特色": str(src.get("feature", "") or ""),
        "ERP品號": str(src.get("erp_sku", "") or ""),
        "國際條碼": str(src.get("barcode", "") or ""),
        "價格\n進價/成本價": str(src.get("cost_price", "") or ""),
        "價格\n建議售價": str(src.get("list_price", "") or ""),
        "商品材積\n長(cm)": str(src.get("length_cm", "") or ""),
        "商品材積\n寬(cm)": str(src.get("width_cm", "") or ""),
        "商品材積\n高(cm)": str(src.get("height_cm", "") or ""),
        "重量(KG)": str(src.get("weight_kg", "") or ""),
        "效期天數": str(src.get("warranty_days", "") or ""),
        "is_shelflife": str(src.get("is_shelflife", "") or ""),
        "販售類型": str(src.get("selling_type", "") or ""),
        "是否原箱販售": str(src.get("is_box_sale", "") or ""),
        "規格名稱1": str(src.get("spec_name_1", "") or ""),
        "規格內容1": str(src.get("spec_value_1", "") or ""),
        "規格名稱2": str(src.get("spec_name_2", "") or ""),
        "規格內容2": str(src.get("spec_value_2", "") or ""),
        "商品分類": str(src.get("category", "") or ""),
    }


def iter_rows_as_dict(ws: Worksheet, header_row: int, start_row: int) -> List[Dict[str, str]]:
    headers: List[str] = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        headers.append("" if v is None else str(v).strip())

    rows: List[Dict[str, str]] = []
    for r in range(start_row, ws.max_row + 1):
        item: Dict[str, str] = {}
        empty = True
        for c, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(r, c).value
            s = "" if v is None else str(v).strip()
            if s:
                empty = False
            item[h] = s
        if not empty:
            rows.append(item)
    return rows


def compose_description(item: Dict[str, str]) -> str:
    parts: List[str] = []
    slogan = item.get("商品特色標語", "").strip()
    feats = item.get("\n商品特色", "").strip()  # 你原本的欄位名
    if slogan:
        parts.append(slogan)
    if feats:
        parts.append(feats)
    return "\n".join([p for p in parts if p]).strip()


# -----------------------------
# Template helpers
# -----------------------------
def build_col_index_from_row(ws: Worksheet, header_row: int) -> Dict[str, int]:
    idx: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is None:
            continue
        name = str(v).strip()
        if name and name not in idx:
            idx[name] = c
    return idx


def read_template_headers(ws: Worksheet, header_row: int) -> List[str]:
    headers: List[str] = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        headers.append("" if v is None else str(v).strip())
    while headers and headers[-1] == "":
        headers.pop()
    return headers


# -----------------------------
# Category search
# -----------------------------
@dataclass(frozen=True)
class CategoryRow:
    cat_id: str
    path: str
    leaf: str


def _norm(s: str) -> str:
    s = s.strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def load_categories(category_xls: Path) -> List[CategoryRow]:
    try:
        df = pd.read_excel(category_xls)
    except ImportError as e:
        raise RuntimeError("Reading .xls requires xlrd. Please run: pip install xlrd") from e

    def col(name: str) -> str:
        for c in df.columns:
            if str(c).strip() == name:
                return c
        raise KeyError(f"column not found: {name}. available={list(df.columns)}")

    c1 = col("Level 1 Category")
    c2 = col("Level 2 Category")
    c3 = col("Level 3 Category")
    c4 = col("Level 4 Category")
    c5 = col("Level 5 Category")
    cid = col("Category ID")

    out: List[CategoryRow] = []
    for _, row in df.iterrows():
        levels: List[str] = []
        for key in (c1, c2, c3, c4, c5):
            v = row.get(key)
            if v is None:
                continue
            if isinstance(v, float) and math.isnan(v):
                continue
            s = str(v).strip()
            if s:
                levels.append(s)

        if not levels:
            continue

        v_id = row.get(cid)
        if v_id is None or (isinstance(v_id, float) and math.isnan(v_id)):
            continue

        if isinstance(v_id, float) and float(v_id).is_integer():
            cat_id = str(int(v_id))
        else:
            cat_id = str(v_id).strip()

        path = " > ".join(levels)
        leaf = levels[-1]
        out.append(CategoryRow(cat_id=cat_id, path=path, leaf=leaf))

    return out


def rank_categories(categories: List[CategoryRow], query: str, topk: int = 5) -> List[Tuple[int, int, CategoryRow]]:
    q = _norm(query)
    if not q:
        return []

    q_tokens = [t for t in re.split(r"[\s/,_\-]+", q) if t]

    scored: List[Tuple[int, int, CategoryRow]] = []
    for r in categories:
        leaf = _norm(r.leaf)
        path = _norm(r.path)

        score = 0
        if leaf == q:
            score = 300
        elif q in leaf or leaf in q:
            score = 220
        elif any(tok and tok in leaf for tok in q_tokens):
            score = 180
        elif q in path:
            score = 120
        else:
            continue

        scored.append((score, len(path), r))

    scored.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return scored[:topk]


def find_best_category_id(categories: List[CategoryRow], query: str) -> Optional[str]:
    ranked = rank_categories(categories, query, topk=1)
    return ranked[0][2].cat_id if ranked else None


# -----------------------------
# Output helpers
# -----------------------------
def to_int_str(v: Any, *, default: Optional[int] = None) -> str:
    s = "" if v is None else str(v).strip()
    if not s:
        return "" if default is None else str(int(default))
    s = s.replace(",", "")
    try:
        f = float(s)
        return str(int(round(f)))
    except Exception:
        return "" if default is None else str(int(default))


def to_decimal_str(v: Any, *, default: Optional[float] = None, ndigits: int = 2) -> str:
    s = "" if v is None else str(v).strip()
    if not s:
        return "" if default is None else f"{float(default):.{ndigits}f}"
    s = s.replace(",", "")
    try:
        f = float(s)
        return f"{f:.{ndigits}f}"
    except Exception:
        return "" if default is None else f"{float(default):.{ndigits}f}"


def to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except Exception:
        return None


def normalize_yes_no(v: Any, *, default: str = "No") -> str:
    s = "" if v is None else str(v).strip().lower()
    if not s:
        return default
    if s in {"yes", "y", "true", "1", "是"}:
        return "Yes"
    if s in {"no", "n", "false", "0", "否"}:
        return "No"
    return default


def clamp_name(name: str, *, min_len: int = 10, max_len: int = 60, fallback: str = "未命名商品") -> str:
    s = (name or "").strip()
    if not s:
        s = fallback
    if len(s) < min_len:
        s = (s + " " + fallback).strip()
    if len(s) > max_len:
        s = s[:max_len]
    return s


def write_clean_output(headers: List[str], data_rows: List[List[str]], output_xlsx: Path) -> None:
    """
    只輸出：
    - 第 1 列：欄位列
    - 第 2 列起：資料列
    不輸出「必填/選填」那一列
    """
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    assert out_ws is not None
    out_ws.title = "Template"

    ncols = len(headers)

    # 1) 欄位列
    out_ws.append(headers[:ncols])

    # 2) 資料列
    for row in data_rows:
        if len(row) < ncols:
            row = row + [""] * (ncols - len(row))
        out_ws.append(row[:ncols])

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output_xlsx)


def init_output_workbook(headers: List[str]) -> tuple[Workbook, Worksheet]:
    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    assert out_ws is not None
    out_ws.title = "Template"
    out_ws.append(headers)
    return out_wb, out_ws


def append_row_and_flush(out_wb: Workbook, out_ws: Worksheet, row: List[str], ncols: int, output_xlsx: Path) -> None:
    if len(row) < ncols:
        row = row + [""] * (ncols - len(row))
    out_ws.append(row[:ncols])
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output_xlsx)


def write_error_report(errors: List[Dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(errors)
    df.to_excel(path, index=False)


# -----------------------------
# Main
# -----------------------------
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input-xlsx", type=Path, default=Path("test.xlsx"), help="來源 test.xlsx（use-gsheet 時可忽略）")
    ap.add_argument("--template", type=Path, required=True, help="Shopee 大量上架範本（.xlsm/.xlsx）")
    ap.add_argument("--category-xls", type=Path, required=True, help="Shopee 類別目錄（Shopee_category_list.xls）")
    ap.add_argument("--output-xlsx", type=Path, required=True, help="輸出 .xlsx（只保留欄位列+資料列）")
    ap.add_argument("--error-xlsx", type=Path, default=Path("out/upload_errors.xlsx"), help="錯誤報表輸出（預設 out/upload_errors.xlsx）")

    ap.add_argument("--input-sheet", type=str, default="", help="來源工作表名稱（留空=第一張）")
    ap.add_argument("--template-sheet", type=str, default="Template", help="範本工作表（預設=Template）")
    ap.add_argument("--input-header-row", type=int, default=1, help="來源表頭列（預設=1）")
    ap.add_argument("--skip-row-2", action="store_true", help="略過第二列（欄位說明）")

    ap.add_argument("--category-query-col", type=str, default="商品分類", help="來源用來搜尋類別的欄位名（若不存在就用 default-category-query）")
    ap.add_argument("--default-category-query", type=str, default="吹風機家庭用", help="預設類別搜尋關鍵字")
    ap.add_argument("--fail-on-category-miss", action="store_true", help="找不到類別就直接把該列列為錯誤（不寫入輸出）")
    ap.add_argument("--print-category-topk", type=int, default=0, help=">0 則印出每列類別搜尋 TopK 候選（debug 用）")
    ap.add_argument("--use-rag-ai-category", action="store_true", help="使用 RAG + LLM 重排做類別判斷（參考 momo）")
    ap.add_argument("--category-top-k", type=int, default=10, help="RAG 候選數量（use-rag-ai-category 時生效）")
    ap.add_argument("--min-category-confidence", type=float, default=0.55, help="最小 final_confidence 門檻")
    ap.add_argument("--min-category-retrieval-score", type=float, default=0.12, help="最小 top1 retrieval 分數門檻")
    ap.add_argument("--min-category-retrieval-gap", type=float, default=0.005, help="最小 top1-top2 retrieval gap 門檻")
    ap.add_argument("--category-force-rebuild-index", action="store_true", help="強制重建 category embedding cache")
    ap.add_argument("--auto-start-ollama", dest="auto_start_ollama", action="store_true", help="需要時自動啟動 Ollama（預設開啟）")
    ap.add_argument("--no-auto-start-ollama", dest="auto_start_ollama", action="store_false", help="停用自動啟動 Ollama")
    ap.set_defaults(auto_start_ollama=True)
    ap.add_argument("--ollama-ready-timeout", type=float, default=90.0, help="等待 Ollama 就緒秒數")
    ap.add_argument("--use-rag-result-memory", dest="use_rag_result_memory", action="store_true", help="啟用 RAG 分類結果記憶庫（預設開啟）")
    ap.add_argument("--no-rag-result-memory", dest="use_rag_result_memory", action="store_false", help="停用 RAG 分類結果記憶庫")
    ap.set_defaults(use_rag_result_memory=True)
    ap.add_argument("--reuse-needs-review-memory", action="store_true", help="允許重用 needs_review 的歷史結果（預設只重用穩定結果）")

    ap.add_argument("--require-dim", action="store_true", help="強制長寬高必填且為整數（不符合列入錯誤）")
    ap.add_argument("--require-weight", action="store_true", help="強制重量必填且 >0（不符合列入錯誤）")
    ap.add_argument("--use-gsheet", action="store_true", help="改用 done/shared/config/gsheet_pm_source.json 讀取 Google Sheet")

    args = ap.parse_args()

    # 1) source
    if args.use_gsheet:
        gs = load_gsheet_pm_source_config()
        src_rows_raw = read_sheet_records(
            credentials_path=str(gs["cred_abs_path"]),
            spreadsheet_id=str(gs.get("spreadsheet_id", "") or ""),
            worksheet_name=str(gs.get("worksheet_name", "") or ""),
            header_row=int(gs.get("header_row") or 2),
            data_start_row=int(gs.get("data_start_row") or 4),
        )
        src_rows = [gsheet_row_to_shopee_row(x) for x in src_rows_raw]
        if not src_rows:
            raise ValueError("Google Sheet has no data rows.")
    else:
        src_wb: Workbook = openpyxl.load_workbook(args.input_xlsx, data_only=True)
        src_ws: Worksheet = src_wb[args.input_sheet] if args.input_sheet else src_wb.worksheets[0]
        input_start_row = 3 if args.skip_row_2 else (args.input_header_row + 1)
        src_rows = iter_rows_as_dict(src_ws, header_row=args.input_header_row, start_row=input_start_row)
        if not src_rows:
            raise ValueError("source xlsx has no data rows (after skipping).")

    # 2) categories
    categories = load_categories(args.category_xls)
    if args.use_rag_ai_category:
        preload_index(
            args.category_xls,
            force_rebuild=args.category_force_rebuild_index,
            auto_start_ollama=args.auto_start_ollama,
            ollama_ready_timeout_sec=args.ollama_ready_timeout,
        )

    # 3) template (只讀欄位列，不再讀「必填/選填」列)
    tpl_wb: Workbook = openpyxl.load_workbook(args.template, data_only=True)
    if args.template_sheet not in tpl_wb.sheetnames:
        raise ValueError(f"template sheet not found: {args.template_sheet}. available={tpl_wb.sheetnames}")
    tpl_ws: Worksheet = tpl_wb[args.template_sheet]

    zh_header_row = 3  # template 欄位列
    tpl_headers = read_template_headers(tpl_ws, zh_header_row)
    tpl_col_idx = build_col_index_from_row(tpl_ws, zh_header_row)
    shelf_life_flag_aliases = {"效期品", "is_shelflife", "Is Shelf Life", "是否有效期", "是否效期"}
    if not any(h in shelf_life_flag_aliases for h in tpl_headers):
        # 有些舊版模板沒有「效期品」欄，為了與新版上傳格式相容，動態補一欄。
        # 順序需為：重量 -> 效期品 -> 效期天數 -> 販售類型。
        if "重量" in tpl_headers:
            insert_at = tpl_headers.index("重量") + 1
        elif "效期天數" in tpl_headers:
            insert_at = tpl_headers.index("效期天數")
        else:
            insert_at = len(tpl_headers)
        tpl_headers.insert(insert_at, "效期品")
        tpl_col_idx = {name: i + 1 for i, name in enumerate(tpl_headers)}
    ncols = len(tpl_headers)

    def set_val(row_arr: List[str], field_name: str, value: str) -> None:
        col = tpl_col_idx.get(field_name)
        if not col:
            return
        if 1 <= col <= ncols:
            row_arr[col - 1] = value

    def set_val_candidates(row_arr: List[str], field_names: List[str], value: str) -> None:
        for n in field_names:
            set_val(row_arr, n, value)

    out_wb, out_ws = init_output_workbook(tpl_headers)
    processed_rows = 0
    errors: List[Dict[str, Any]] = []

    for i, item in enumerate(src_rows, start=1):
        row_arr = [""] * ncols
        row_errors: List[str] = []

        # category query
        q = item.get(args.category_query_col, "").strip()
        if not q:
            q = args.default_category_query.strip()

        cat_id = ""
        cat_meta: Dict[str, Any] = {}
        if args.use_rag_ai_category:
            fallback_id = find_best_category_id(categories, q)
            rag_result = predict_category(
                {
                    "category_query": q,
                    "name": item.get("銷售品名(相同當作同一賣場)", ""),
                    "slogan": item.get("商品特色標語", ""),
                    "feature": item.get("\n商品特色", ""),
                    "description": compose_description(item),
                    "brand": item.get("品牌", ""),
                    "erp_sku": item.get("ERP品號", ""),
                    "spec_1": f"{item.get('規格名稱1', '')}={item.get('規格內容1', '')}",
                    "spec_2": f"{item.get('規格名稱2', '')}={item.get('規格內容2', '')}",
                },
                category_xls=args.category_xls,
                top_k=args.category_top_k,
                min_confidence=args.min_category_confidence,
                min_retrieval_score=args.min_category_retrieval_score,
                min_retrieval_gap=args.min_category_retrieval_gap,
                auto_start_ollama=args.auto_start_ollama,
                ollama_ready_timeout_sec=args.ollama_ready_timeout,
                use_result_memory=args.use_rag_result_memory,
                reuse_needs_review_memory=args.reuse_needs_review_memory,
            )
            final_result = rag_result.get("final_result", {})
            cat_id = str(final_result.get("category_code", "") or "")
            cat_meta = {
                "category_path": final_result.get("category_path", ""),
                "final_confidence": final_result.get("final_confidence", 0.0),
                "top1_retrieval_score": final_result.get("top1_retrieval_score", 0.0),
                "retrieval_gap_12": final_result.get("retrieval_gap_12", 0.0),
                "needs_review": final_result.get("needs_review", False),
                "review_reasons": "|".join(final_result.get("review_reasons", [])),
                "memory_hit": rag_result.get("memory_hit", False),
            }
            if args.print_category_topk:
                cands = rag_result.get("top_k_candidates", [])[: args.print_category_topk]
                print(f"[row {i}] category query='{q}' rag top{args.print_category_topk}:")
                for c in cands:
                    print(
                        f"  - id={c.get('cat_id', '')} score={float(c.get('retrieval_score', 0.0) or 0.0):.4f} "
                        f"path={c.get('path', '')}"
                    )
            if not cat_id and fallback_id:
                cat_id = fallback_id
                row_errors.append("RAG/LLM category empty, fallback to keyword match")
            if bool(final_result.get("needs_review", False)):
                row_errors.append(
                    "category needs review: "
                    + "|".join(final_result.get("review_reasons", []))
                )
        else:
            if args.print_category_topk and q:
                ranked = rank_categories(categories, q, topk=args.print_category_topk)
                print(f"[row {i}] category query='{q}' top{args.print_category_topk}:")
                for sc, ln, r in ranked:
                    print(f"  - id={r.cat_id} score={sc} path={r.path}")
            cat_id = find_best_category_id(categories, q) or ""

        if not cat_id:
            row_errors.append(f"category not found: query='{q}'")
        else:
            set_val(row_arr, "類別", to_int_str(cat_id))

        # 商品名稱：10~60字
        name = item.get("銷售品名(相同當作同一賣場)", "").strip()
        set_val(row_arr, "商品名稱 (品)", clamp_name(name))

        # 商品描述
        desc = compose_description(item)
        if desc:
            set_val(row_arr, "商品描述", desc)

        # ERP品號 -> 主商品貨號 + 供應商料號（必填之一）
        erp_sku = item.get("ERP品號", "").strip()
        if erp_sku:
            set_val(row_arr, "主商品貨號", erp_sku)
            set_val(row_arr, "供應商料號", erp_sku)

        # 國際條碼：空值 => NO EAN/UPC，且不能有底線
        barcode = item.get("國際條碼", "").strip()
        if not barcode:
            barcode = "NO EAN/UPC"
        if "_" in barcode:
            barcode = barcode.replace("_", "")
        set_val(row_arr, "國際條碼", barcode)

        # 價格：稅後進價 / 市價
        cost_raw = item.get("價格\n進價/成本價", "")
        orig_raw = item.get("價格\n建議售價", "")
        cost_val = to_float(cost_raw)
        orig_val = to_float(orig_raw)

        if cost_val is None:
            row_errors.append("稅後進價為空或非數字")
        elif cost_val <= 0:
            row_errors.append("稅後進價不可 <= 0")

        if orig_val is None:
            row_errors.append("市價為空或非數字")
        elif orig_val <= 0:
            row_errors.append("市價不可 <= 0")

        if cost_val is not None and orig_val is not None and cost_val >= orig_val:
            row_errors.append("稅後進價不可 >= 市價")

        if cost_val is not None:
            set_val(row_arr, "稅後進價", to_decimal_str(cost_val, ndigits=2))
        if orig_val is not None:
            set_val(row_arr, "市價", to_decimal_str(orig_val, ndigits=2))

        # 長寬高
        length_s = to_int_str(item.get("商品材積\n長(cm)", ""))
        width_s = to_int_str(item.get("商品材積\n寬(cm)", ""))
        height_s = to_int_str(item.get("商品材積\n高(cm)", ""))

        set_val(row_arr, "長度", length_s)
        set_val(row_arr, "寬度", width_s)
        set_val(row_arr, "高度", height_s)

        if args.require_dim:
            if not length_s or not width_s or not height_s:
                row_errors.append("長寬高必填（require-dim）")

        # 重量
        weight_val = to_float(item.get("重量(KG)", ""))
        if weight_val is not None:
            weight_s = to_decimal_str(weight_val, ndigits=2)
            set_val_candidates(row_arr, ["重量", "Net weight"], weight_s)
        if args.require_weight:
            if weight_val is None or weight_val <= 0:
                row_errors.append("重量必填且需 >0（require-weight）")

        # 效期天數：預設 365
        shelf_life_days = to_int_str(item.get("效期天數", ""), default=365)
        set_val_candidates(row_arr, ["效期天數", "Shelf Life(Days)"], shelf_life_days)

        # is_shelflife：Yes/No
        src_is_shelflife = item.get("is_shelflife", "").strip()
        if src_is_shelflife:
            is_shelflife = normalize_yes_no(src_is_shelflife, default="No")
        else:
            is_shelflife = "Yes" if (to_float(shelf_life_days) or 0) > 0 else "No"
        set_val_candidates(
            row_arr,
            ["is_shelflife", "Is Shelf Life", "是否有效期", "是否效期", "效期品"],
            is_shelflife,
        )

        # 販售類型：預設 2=散賣(個)
        selling_type = to_int_str(item.get("販售類型", ""), default=2)
        set_val_candidates(row_arr, ["販售類型", "Selling type"], selling_type)

        # 是否原箱販售：是/否（預設 否）
        v_case = item.get("是否原箱販售", "").strip()
        if v_case not in ("是", "否"):
            v_case = "否"
        set_val_candidates(row_arr, ["是否原箱販售", "is_box_sale"], v_case)

        # 規格
        for src_h, dst_field in {
            "規格名稱1": "規格名稱 1",
            "規格內容1": "規格選項 1",
            "規格名稱2": "規格名稱 2",
            "規格內容2": "規格選項 2",
        }.items():
            v = item.get(src_h, "").strip()
            if v:
                set_val(row_arr, dst_field, v)

        if row_errors:
            err = {
                "row_index": i,
                "category_query": q,
                "category_path": cat_meta.get("category_path", ""),
                "category_final_confidence": cat_meta.get("final_confidence", ""),
                "category_top1_retrieval_score": cat_meta.get("top1_retrieval_score", ""),
                "category_retrieval_gap_12": cat_meta.get("retrieval_gap_12", ""),
                "category_needs_review": cat_meta.get("needs_review", ""),
                "category_review_reasons": cat_meta.get("review_reasons", ""),
                    "category_memory_hit": cat_meta.get("memory_hit", ""),
                "erp_sku": erp_sku,
                "name": name,
                "errors": " | ".join(row_errors),
            }
            errors.append(err)
            # 每處理完一列就更新錯誤報表（即時落盤）
            write_error_report(errors, args.error_xlsx)

            if args.fail_on_category_miss and any(e.startswith("category not found") for e in row_errors):
                continue

        append_row_and_flush(out_wb, out_ws, row_arr, ncols, args.output_xlsx)
        processed_rows += 1

    print(f"[OK] wrote: {args.output_xlsx} rows={processed_rows}")

    if errors:
        write_error_report(errors, args.error_xlsx)
        print(f"[WARN] wrote error report: {args.error_xlsx} rows={len(errors)}")
    else:
        print("[OK] no validation errors")


if __name__ == "__main__":
    main()
