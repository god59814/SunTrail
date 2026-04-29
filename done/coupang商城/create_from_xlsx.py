import argparse
import csv
import hmac
import hashlib
import json
import sys
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

_SHARED_LIB = Path(__file__).resolve().parent.parent / "shared"
if str(_SHARED_LIB) not in sys.path:
    sys.path.insert(0, str(_SHARED_LIB))
from work_paths import load_gsheet_pm_source, resolve_cred_path, work_root_from

import httpx
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from gsheet_reader import read_sheet_records

from local_image_vendorpath import LocalImageHosting
from category_search import resolve_category_code
from shipping_fee import fee_from_dimension_sum, get_dimension_sum
from ollama_test.llm_category_service import predict_category, preload_index

# ========= paths =========
CONFIG_PATH = Path("config/login_info.json")
MAPPING_PATH = Path("mapping.json")
OUT_DIR = Path("out")
OUT_DIR.mkdir(parents=True, exist_ok=True)
REPORT_CSV = OUT_DIR / "upload_report.csv"


def _parse_display_category_code(raw: Any) -> Optional[int]:
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    digits = "".join(ch for ch in s if ch.isdigit())
    if not digits:
        return None
    try:
        return int(digits)
    except ValueError:
        return None


def _category_report_fields(
    *,
    resolved_mode: str,
    llm_final: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    lf = llm_final or {}
    return {
        "category_mode": resolved_mode,
        "llm_category_code": str(lf.get("category_code", "")),
        "llm_confidence": str(lf.get("confidence", "")),
        "llm_reason": str(lf.get("reason", "")),
    }


def build_product_for_llm(row: Dict[str, Any]) -> Dict[str, str]:
    return {
        "sale_product_name": str(get_cell(row, "sale_product_name", "") or "").strip(),
        "erp_product_name": str(get_cell(row, "erp_product_name", "") or "").strip(),
        "brand": str(get_cell(row, "brand", "") or "").strip(),
        "slogan": str(get_cell(row, "slogan", "") or "").strip(),
        "feature": str(get_cell(row, "feature", "") or "").strip(),
        "product_description": str(
            get_cell(row, "product_description", "") or ""
        ).strip(),
        "model_no": str(get_cell(row, "model_no", "") or "").strip(),
        "spec_name_1": str(get_cell(row, "spec_name_1", "") or "").strip(),
        "spec_value_1": str(get_cell(row, "spec_value_1", "") or "").strip(),
        "spec_name_2": str(get_cell(row, "spec_name_2", "") or "").strip(),
        "spec_value_2": str(get_cell(row, "spec_value_2", "") or "").strip(),
    }


def _merge_gsheet_config(gs_from_mapping: Dict[str, Any]) -> Dict[str, Any]:
    wr = work_root_from(Path(__file__))
    base = dict(load_gsheet_pm_source(wr))
    over = dict(gs_from_mapping or {})
    merged = {**base, **over}
    cred = str(merged.get("cred_path", "") or "").strip()
    if cred:
        merged["cred_path"] = resolve_cred_path(wr, cred)
    return merged


# Product Creation API
API_PATH = "/v2/providers/seller_api/apis/api/v1/marketplace/seller-products"


# ========= config =========
@dataclass(frozen=True)
class CoupangConfig:
    access_key: str
    secret_key: str
    vendor_id: str
    base_url: str


def load_config() -> CoupangConfig:
    obj = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    access_key = str(obj.get("access_key", "")).strip()
    secret_key = str(obj.get("secret_key", "")).strip()
    vendor_id = str(obj.get("vendor_id", "")).strip()
    base_url = str(obj.get("base_url", "https://api-gateway.coupang.com")).strip()

    if not access_key or not secret_key:
        raise ValueError("config/login_info.json missing access_key or secret_key")
    if not vendor_id:
        raise ValueError("config/login_info.json missing vendor_id")

    return CoupangConfig(access_key, secret_key, vendor_id, base_url)


def load_mapping() -> Dict[str, Any]:
    if not MAPPING_PATH.exists():
        raise FileNotFoundError(f"Missing {MAPPING_PATH.resolve()}")
    return json.loads(MAPPING_PATH.read_text(encoding="utf-8"))


def utc_signed_date() -> str:
    # 你目前在 TW API 上用 yymmddTHHMMSSZ 跑通 GET，所以沿用
    return time.strftime("%y%m%dT%H%M%SZ", time.gmtime())


def to_int(v: Any, default: int = 0) -> int:
    if v is None:
        return default
    if isinstance(v, (int,)):
        return int(v)
    if isinstance(v, float):
        return int(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return default
    try:
        return int(float(s))
    except Exception:
        return default


def build_authorization(
    method: str, path: str, query_str: str, access_key: str, secret_key: str
) -> str:
    signed_date = utc_signed_date()
    message = f"{signed_date}{method.upper()}{path}{query_str}"

    signature = hmac.new(
        key=secret_key.encode("utf-8"),
        msg=message.encode("utf-8"),
        digestmod=hashlib.sha256,
    ).hexdigest()

    return (
        "CEA algorithm=HmacSHA256, "
        f"access-key={access_key}, "
        f"signed-date={signed_date}, "
        f"signature={signature}"
    )


def http_post(
    cfg: CoupangConfig,
    market: str,
    path: str,
    body: Dict[str, Any],
    timeout: float = 60.0,
) -> Tuple[int, str, Dict[str, str], Dict[str, str]]:
    method = "POST"
    query_str = ""
    auth = build_authorization(method, path, query_str, cfg.access_key, cfg.secret_key)

    req_headers = {
        "Authorization": auth,
        "Content-Type": "application/json;charset=UTF-8",
        "Accept": "application/json",
        "X-MARKET": market,
        "X-Requested-By": cfg.vendor_id,
    }
    url = f"{cfg.base_url}{path}"

    t = httpx.Timeout(timeout, connect=10.0, read=timeout, write=timeout, pool=10.0)

    try:
        with httpx.Client(timeout=t, follow_redirects=True) as client:
            r = client.post(url, headers=req_headers, json=body)
            return r.status_code, r.text, dict(r.headers), dict(req_headers)
    except Exception as e:
        err = {
            "error": str(e),
            "traceback": traceback.format_exc(),
            "url": url,
        }
        return 599, json.dumps(err, ensure_ascii=False, indent=2), {}, dict(req_headers)


# ========= xlsx =========
def read_xlsx_rows(
    xlsx_path: Path, skip_rows: List[int], header_row: int
) -> Tuple[List[str], List[Dict[str, Any]]]:
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    if ws is None:
        raise ValueError("Workbook has no active worksheet")
    assert isinstance(ws, Worksheet)

    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    headers = [("" if h is None else str(h).strip()) for h in headers]

    rows: List[Dict[str, Any]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        if r in skip_rows:
            continue

        row_dict: Dict[str, Any] = {}
        empty = True
        for c, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() != "":
                empty = False
            row_dict[h] = v

        if not empty:
            rows.append(row_dict)

    return headers, rows


def get_cell(row: Dict[str, Any], col_name: str, default: Any = "") -> Any:
    return row.get(col_name, default)


def normalize_tax_type(v: Any) -> Optional[str]:
    """
    Excel 可能出現：免稅/應稅、免稅、應稅、FREE、TAX、0/1、True/False... 這裡統一轉成 Coupang 需要的值。
    回傳 "TAX" / "FREE" / None(代表不覆蓋，沿用 item_defaults)
    """
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None

    s_upper = s.upper()

    # 已經是 API 值
    if s_upper in ("TAX", "FREE"):
        return s_upper

    # 常見中文
    if "應稅" in s:
        return "TAX"
    if "免稅" in s:
        return "FREE"

    # 常見 boolean / 數字習慣：這裡做保守解讀
    if s_upper in ("Y", "YES", "TRUE", "1"):
        # 有些人會用 1 表示應稅
        return "TAX"
    if s_upper in ("N", "NO", "FALSE", "0"):
        # 有些人會用 0 表示免稅
        return "FREE"

    return None


def group_rows(
    rows: List[Dict[str, Any]], group_key_col: str, fallback_key_col: str
) -> Dict[str, List[Dict[str, Any]]]:
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        key = str(get_cell(r, group_key_col, "") or "").strip()
        if not key:
            key = str(get_cell(r, fallback_key_col, "") or "").strip() or "UNKNOWN"
        groups.setdefault(key, []).append(r)
    return groups


# ========= report =========
def write_report(rows: List[Dict[str, Any]]) -> None:
    with REPORT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=[
                "group_name",
                "category_code",
                "category_mode",
                "llm_category_code",
                "llm_confidence",
                "llm_reason",
                "payload_file",
                "uploaded",
                "http_status",
                "message",
                "missing",
            ],
        )
        w.writeheader()
        for r in rows:
            w.writerow(r)


# ========= TW normalize =========
def normalize_shipping_fee_tw(payload: Dict[str, Any]) -> None:
    msi = payload.get("multiShippingInfos") or []
    home_ship = None
    if isinstance(msi, list):
        for x in msi:
            if isinstance(x, dict) and x.get("deliveryCompanyType") == "HOME_DELIVERY":
                home_ship = x
                break

    if not isinstance(home_ship, dict):
        return

    payload["deliveryMethod"] = home_ship.get("deliveryMethod", "SEQUENCIAL")
    payload["deliveryCompanyCode"] = home_ship.get("deliveryCompanyCode")
    payload["deliveryChargeType"] = home_ship.get("deliveryChargeType")
    payload["deliveryCharge"] = home_ship.get("deliveryCharge", 0)
    payload["freeShipOverAmount"] = home_ship.get("freeShipOverAmount", 0)
    payload["deliveryChargeOnReturn"] = home_ship.get("deliveryChargeOnReturn", 0)

    if home_ship.get("outboundShippingPlaceId") is not None:
        payload["outboundShippingPlaceCode"] = home_ship["outboundShippingPlaceId"]

    if (
        payload.get("deliveryChargeType") == "FREE"
        and payload.get("deliveryCharge") is None
    ):
        payload["deliveryCharge"] = 0
    if payload.get("freeShipOverAmount") is None:
        payload["freeShipOverAmount"] = 0


def apply_delivery_charge_override(
    payload: Dict[str, Any], delivery_charge: int
) -> None:
    """
    以 CLI/UI 指定的運費金額覆寫 payload 的 deliveryCharge（對應長+寬+高級距）。
    同時更新 multiShippingInfos 中 HOME_DELIVERY 的 deliveryCharge。
    """
    payload["deliveryCharge"] = delivery_charge
    msi = payload.get("multiShippingInfos") or []
    if isinstance(msi, list):
        for x in msi:
            if isinstance(x, dict) and x.get("deliveryCompanyType") == "HOME_DELIVERY":
                x["deliveryCharge"] = delivery_charge
                break


def apply_shipping_mode_override(
    payload: Dict[str, Any],
    delivery_charge_type: str,
    free_over_amount: Optional[int],
) -> None:
    """
    依 UI/CLI 指定的運費模式覆寫：
    - deliveryChargeType: NOT_FREE / CONDITIONAL_FREE / FREE
    - freeShipOverAmount: 滿額免運門檻（僅 CONDITIONAL_FREE 使用）
    其他欄位（deliveryCharge）已由 apply_delivery_charge_override / 尺寸自動運費決定。
    """
    t = (delivery_charge_type or "").strip().upper()
    if t not in ("NOT_FREE", "CONDITIONAL_FREE", "FREE"):
        return

    payload["deliveryChargeType"] = t

    if t == "FREE":
        payload["freeShipOverAmount"] = 0
    elif t == "CONDITIONAL_FREE":
        amt = int(free_over_amount or 0)
        if amt <= 0:
            # 若未給有效門檻，就不覆寫，以免產生不合法設定
            return
        payload["freeShipOverAmount"] = amt
    else:  # NOT_FREE
        payload["freeShipOverAmount"] = 0

    # 同步 multiShippingInfos
    msi = payload.get("multiShippingInfos") or []
    if isinstance(msi, list):
        for x in msi:
            if isinstance(x, dict) and x.get("deliveryCompanyType") == "HOME_DELIVERY":
                x["deliveryChargeType"] = payload["deliveryChargeType"]
                x["freeShipOverAmount"] = payload.get("freeShipOverAmount", 0)
                break


def apply_return_charge_override(payload: Dict[str, Any], return_charge: int) -> None:
    """覆寫退貨運費：payload 頂層與 multiReturnInfos / multiShippingInfos 中的 returnCharge、deliveryChargeOnReturn。"""
    payload["returnCharge"] = return_charge
    payload["deliveryChargeOnReturn"] = return_charge
    for mri in payload.get("multiReturnInfos") or []:
        if isinstance(mri, dict):
            mri["returnCharge"] = return_charge
    for msi in payload.get("multiShippingInfos") or []:
        if isinstance(msi, dict):
            msi["deliveryChargeOnReturn"] = return_charge


def normalize_tw_payload(payload: Dict[str, Any]) -> None:
    if not str(payload.get("registrationType", "")).strip():
        payload["registrationType"] = "NORMAL"

    msi = payload.get("multiShippingInfos") or []
    home_ship = None
    if isinstance(msi, list):
        for x in msi:
            if isinstance(x, dict) and x.get("deliveryCompanyType") == "HOME_DELIVERY":
                home_ship = x
                break

    if isinstance(home_ship, dict):
        payload.setdefault(
            "deliveryMethod", home_ship.get("deliveryMethod", "SEQUENCIAL")
        )
        payload.setdefault("deliveryCompanyCode", home_ship.get("deliveryCompanyCode"))
        payload.setdefault("deliveryChargeType", home_ship.get("deliveryChargeType"))
        payload.setdefault("deliveryCharge", home_ship.get("deliveryCharge"))
        payload.setdefault("freeShipOverAmount", home_ship.get("freeShipOverAmount", 0))
        payload.setdefault(
            "deliveryChargeOnReturn", home_ship.get("deliveryChargeOnReturn")
        )

        ospid = home_ship.get("outboundShippingPlaceId")
        if ospid is not None:
            payload.setdefault("outboundShippingPlaceCode", ospid)

    if (
        payload.get("deliveryChargeType") == "FREE"
        and payload.get("deliveryCharge") is None
    ):
        payload["deliveryCharge"] = 0
    if payload.get("freeShipOverAmount") is None:
        payload["freeShipOverAmount"] = 0

    mri = payload.get("multiReturnInfos") or []
    home_return = None
    if isinstance(mri, list):
        for x in mri:
            if isinstance(x, dict) and x.get("pickUpBranchType") == "HOME":
                home_return = x
                break

    if isinstance(home_return, dict):
        rcc = str(home_return.get("returnCenterCode", "")).strip()
        if rcc:
            home_return["returnCenterCode"] = rcc
            payload["returnCenterCode"] = rcc

        name_val = str(
            home_return.get("returnChargeName")
            or home_return.get("returnChangeName")
            or payload.get("returnChargeName")
            or payload.get("returnChangeName")
            or ""
        ).strip()
        if name_val:
            home_return["returnChargeName"] = name_val
            home_return["returnChangeName"] = name_val
            payload["returnChargeName"] = name_val
            payload["returnChangeName"] = name_val

        for k in [
            "companyContactNumber",
            "returnZipCode",
            "returnAddress",
            "returnAddressDetail",
            "returnCharge",
        ]:
            if k in home_return:
                payload[k] = home_return[k]


def ensure_tw_validation(
    payload: Dict[str, Any], validation: Dict[str, Any]
) -> List[str]:
    missing: List[str] = []

    if bool(validation.get("require_home_delivery", False)):
        infos = payload.get("multiShippingInfos") or []
        types = [x.get("deliveryCompanyType") for x in infos if isinstance(x, dict)]
        if "HOME_DELIVERY" not in types:
            missing.append("multiShippingInfos must include HOME_DELIVERY")

    if bool(validation.get("require_home_return", False)):
        infos = payload.get("multiReturnInfos") or []
        types = [x.get("pickUpBranchType") for x in infos if isinstance(x, dict)]
        if "HOME" not in types:
            missing.append("multiReturnInfos must include HOME")

        for i, x in enumerate(infos):
            if not isinstance(x, dict):
                continue
            rcc = x.get("returnCenterCode", None)
            if rcc is None:
                missing.append(f"multiReturnInfos[{i}].returnCenterCode missing")
                continue
            rcc_s = str(rcc).strip()
            if (not rcc_s) or (rcc_s == "NO_RETURN_CENTERCODE"):
                missing.append(
                    f"multiReturnInfos[{i}].returnCenterCode missing/placeholder"
                )
                continue
            if not rcc_s.isdigit():
                missing.append(
                    f"multiReturnInfos[{i}].returnCenterCode not numeric: {rcc_s}"
                )

    if not str(payload.get("registrationType", "")).strip():
        missing.append("registrationType")
    if payload.get("deliveryChargeType") and payload.get("deliveryCharge") is None:
        missing.append("deliveryCharge")
    if not str(payload.get("returnCenterCode", "")).strip():
        missing.append("returnCenterCode (root mirror)")

    return missing


# ========= payload building =========
IMAGE_CONFIG_PATH = Path("image_config.json")


def _collect_sku_images(
    folder_path: Path,
    sku: str,
) -> Tuple[List[Path], List[Path], List[Path]]:
    """
    根據「SKU 專屬優先，其次 common」的規則，收集 main / detail / desc 三種圖片。

    資料夾結構預期：
      folder_path/
        common/
          main/
          detail/
          desc/
        <sku>/
          main/
          detail/
          desc/

    規則：
    - main/detail/desc 都是「先看 <sku>/子資料夾，若無或為空，再看 common/子資料夾」
    - 不強制檔名格式，只要在對應子資料夾底下，全部視為有效圖片，並依檔名排序。
    """

    def collect_if_any(subdir: Optional[Path]) -> List[Path]:
        if not subdir or not subdir.exists():
            return []
        return sorted([p for p in subdir.iterdir() if p.is_file()])

    sku_dir: Optional[Path] = folder_path / sku if sku else None
    common_dir: Path = folder_path / "common"

    # main 圖片
    main_files: List[Path] = []
    if sku_dir:
        main_files = collect_if_any(sku_dir / "main")
    if not main_files:
        main_files = collect_if_any(common_dir / "main")

    # detail 圖片
    detail_files: List[Path] = []
    if sku_dir:
        detail_files = collect_if_any(sku_dir / "detail")
    if not detail_files:
        detail_files = collect_if_any(common_dir / "detail")

    # desc 圖片
    desc_files: List[Path] = []
    if sku_dir:
        desc_files = collect_if_any(sku_dir / "desc")
    if not desc_files:
        desc_files = collect_if_any(common_dir / "desc")

    return main_files, detail_files, desc_files


def load_image_config(selected_image_folder: str) -> Dict[str, Any]:
    """
    從 image_config.json 讀取對應 image-folder 的圖片配置：
    {
      "folderA": {
        "main": "a.jpg",
        "details": ["b.jpg", "c.jpg"],
        "contents": ["d.jpg"]
      }
    }
    """
    if not selected_image_folder:
        return {}
    if not IMAGE_CONFIG_PATH.exists():
        return {}
    try:
        cfg_all = json.loads(IMAGE_CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    cfg = cfg_all.get(selected_image_folder) or {}
    if not isinstance(cfg, dict):
        return {}
    return cfg


def build_items(
    group_rows_: List[Dict[str, Any]],
    field_map: Dict[str, str],
    item_defaults: Dict[str, Any],
    validation: Dict[str, Any],
    img_host: Optional[LocalImageHosting],
    fallback_main_image_url: str,
    selected_image_folder: str,
) -> Tuple[List[Dict[str, Any]], List[str]]:

    items: List[Dict[str, Any]] = []
    missing_msgs: List[str] = []

    image_cfg = load_image_config(selected_image_folder)

    for idx, row in enumerate(group_rows_, start=1):

        item_name_col = field_map.get("itemName", "")
        item_name = str(get_cell(row, item_name_col, "") or "").strip() or f"item-{idx}"

        sale_price = to_int(get_cell(row, field_map.get("salePrice", ""), 0))
        original_price = to_int(
            get_cell(row, field_map.get("originalPrice", ""), sale_price)
        )

        barcode = str(get_cell(row, field_map.get("barcode", ""), "") or "").strip()
        sku = str(
            get_cell(row, field_map.get("externalVendorSku", ""), "") or ""
        ).strip()

        stock = to_int(get_cell(row, field_map.get("maximumBuyCount", ""), 0))
        stock = max(0, min(stock, 99999))

        # 規格欄位（例如：顏色、容量等）
        spec_name_col = field_map.get("specName1", "")
        spec_val_col = field_map.get("specValue1", "")
        spec_name = str(get_cell(row, spec_name_col, "") or "").strip()
        spec_val = str(get_cell(row, spec_val_col, "") or "").strip()

        # -----------------------
        # 建立 item 物件（一定要先做）
        # -----------------------
        item: Dict[str, Any] = {}
        item.update(item_defaults)

        item.update(
            {
                "itemName": item_name,
                "salePrice": sale_price,
                "originalPrice": original_price,
                "barcode": barcode,
                "emptyBarcode": False if barcode else True,
                "externalVendorSku": sku,
                "maximumBuyCount": stock,
                "images": [],
                "notices": [],
                "attributes": [],
                "contents": [],
            }
        )

        # 稅別（若 Excel 有帶，優先覆寫）
        tax_col = str(field_map.get("taxType", "") or "").strip()
        if tax_col:
            tax_raw = get_cell(row, tax_col, "")
            tax_raw_s = "" if tax_raw is None else str(tax_raw).strip()
            tax_norm = normalize_tax_type(tax_raw)
            if tax_norm:
                item["taxType"] = tax_norm
            elif tax_raw_s:
                missing_msgs.append(
                    f"invalid taxType '{tax_raw_s}' for item: {item_name}"
                )

        # ======================================================
        # 圖片處理（在此階段 item 已經存在，可以安全操作）
        # ======================================================

        main_img = None
        detail_imgs: List[Path] = []
        desc_imgs: List[Path] = []

        contents_text = ""

        if selected_image_folder and img_host:
            folder_path = Path(img_host.assets_dir) / selected_image_folder
            if not folder_path.exists():
                raise ValueError(f"Image folder not found: {folder_path}")

            # 若有從 UI 設定過 image_config.json，優先使用手動設定
            if image_cfg:
                main_name = str(image_cfg.get("main", "") or "").strip()
                if main_name:
                    main_img = folder_path / main_name

                for n in image_cfg.get("details", []) or []:
                    n_s = str(n or "").strip()
                    if n_s:
                        detail_imgs.append(folder_path / n_s)

                for n in image_cfg.get("contents", []) or []:
                    n_s = str(n or "").strip()
                    if n_s:
                        desc_imgs.append(folder_path / n_s)
                contents_text = str(image_cfg.get("contents_text", "") or "").strip()

            # 若尚未透過 image_config 指定圖片，改用「SKU > common」資料夾結構判斷
            if not main_img and not detail_imgs and not desc_imgs:
                main_files, detail_files, desc_files = _collect_sku_images(
                    folder_path, sku
                )
                if main_files:
                    main_img = main_files[0]
                # 多張主圖：第一張當 REPRESENTATION，其餘視為 DETAIL，優先順序在 detail_ 圖片之前
                # if len(main_files) > 1:
                #     detail_imgs.extend(main_files[1:])
                detail_imgs.extend(detail_files)
                desc_imgs.extend(desc_files)

            # 否則退回舊的自動命名規則：檔名含 main/detail/desc
            if not main_img and not detail_imgs and not desc_imgs:
                files = sorted(folder_path.glob("*"))

                for f in files:
                    name = f.name.lower()
                    if "main" in name:
                        main_img = f
                    elif "detail" in name:
                        detail_imgs.append(f)
                    elif "desc" in name:
                        desc_imgs.append(f)

                if not main_img and files:
                    main_img = files[0]

        if not main_img and fallback_main_image_url:
            main_img = fallback_main_image_url

        if main_img:
            if isinstance(main_img, Path):
                main_url = img_host.to_vendor_url(main_img) if img_host else ""
            else:
                main_url = main_img

            if main_url:
                item["images"].append(
                    {
                        "imageOrder": 0,
                        "imageType": "REPRESENTATION",
                        "vendorPath": main_url,
                    }
                )

                # DETAIL 圖：Coupang 通常最多 9 張；要改數量請改這裡的數字
                for i, img in enumerate(detail_imgs[:9]):
                    if img_host is not None:
                        url = img_host.to_vendor_url(img)
                        item["images"].append(
                            {
                                "imageOrder": i + 1,
                                "imageType": "DETAIL",
                                "vendorPath": url,
                            }
                        )

        else:
            if validation.get("require_each_item_main_image"):
                missing_msgs.append(f"missing main image: {item_name}")

        # ======================================================
        # 商品內容圖：優先支援「第一張圖 + 文字」（IMAGE_TEXT），其餘純圖
        # ======================================================
        if desc_imgs and img_host:
            contents_blocks: List[Dict[str, Any]] = []
            urls: List[str] = []
            # 詳細說明圖（contents）：不限制張數；若要限制可改為 desc_imgs[:N]
            for img in desc_imgs:
                urls.append(img_host.to_vendor_url(img))

            # 若有設定 contents_text，第一張用 IMAGE_TEXT（圖+文），其餘維持 IMAGE
            if urls and contents_text:
                contents_blocks.append(
                    {
                        "contentsType": "IMAGE_TEXT",
                        "contentDetails": [
                            {
                                "content": urls[0],
                                "detailType": "IMAGE",
                            },
                            {
                                "content": contents_text,
                                "detailType": "TEXT",
                            },
                        ],
                    }
                )
                rest = urls[1:]
            else:
                rest = urls

            for url in rest:
                contents_blocks.append(
                    {
                        "contentsType": "IMAGE",
                        "contentDetails": [
                            {
                                "content": url,
                                "detailType": "IMAGE",
                            }
                        ],
                    }
                )

            item["contents"] = contents_blocks

        # 若沒有 desc 圖片內容，改用 mapping.json 的 item_defaults.detail_html（HTML 型態）
        if not item.get("contents"):
            detail_html = str(item_defaults.get("detail_html", "") or "").strip()
            if detail_html:
                item["contents"] = [
                    {
                        "contentsType": "HTML",
                        "contentDetails": [
                            {
                                "content": detail_html,
                                "detailType": "TEXT",
                            }
                        ],
                    }
                ]

        # ======================================================
        # 規格 / 選項屬性（型號、顏色等）
        #   這一段對應 backup 版本，避免 Coupang 視為「沒有規格卻有多個 item」而報重複選項值
        # ======================================================

        # 型號 / 產品編號
        model_col = field_map.get("modelNumber", "")
        model_val = str(get_cell(row, model_col, "") or "").strip()
        if not model_val:
            model_val = sku or barcode
        if model_val:
            item["attributes"].append(
                {"attributeTypeName": "型號/產品編號", "attributeValueName": model_val}
            )
        else:
            missing_msgs.append(f"missing required '型號/產品編號': {item_name}")

        # 顏色：優先用 specName/specValue，其次從品名括號內猜
        color_val = ""
        if spec_name and spec_val:
            sn = spec_name.strip()
            if "顏" in sn:
                sn = "顏色"
            if sn == "顏色":
                color_val = spec_val.strip()

        if not color_val and item_name and "(" in item_name and ")" in item_name:
            color_val = item_name.split("(")[-1].split(")")[0].strip()

        if color_val:
            item["attributes"].append(
                {"attributeTypeName": "顏色", "attributeValueName": color_val}
            )
        else:
            missing_msgs.append(f"missing required '顏色': {item_name}")

        # 其他規格（排除已經用掉的「顏色」與「型號/產品編號」）
        if spec_name and spec_val:
            sn = spec_name.strip()
            if "顏" in sn:
                sn = "顏色"
            if sn not in {"顏色", "型號/產品編號"}:
                item["attributes"].append(
                    {"attributeTypeName": sn, "attributeValueName": spec_val}
                )

        items.append(item)

    return items, missing_msgs


def build_payload(
    cfg: CoupangConfig,
    group_name: str,
    group_rows_: List[Dict[str, Any]],
    category_code: int,
    payload_field: str,
    field_map: Dict[str, str],
    payload_defaults: Dict[str, Any],
    item_defaults: Dict[str, Any],
    validation: Dict[str, Any],
    img_host: Optional[LocalImageHosting],
    fallback_main_image_url: str,
    selected_image_folder: str,
) -> Tuple[Dict[str, Any], List[str]]:
    brand_col = field_map.get("brand", "")
    brand = str(get_cell(group_rows_[0], brand_col, "") or "").strip()

    # 賣場商品名（同一 ERP 多規格分開上架時，品名以 erp_product_name 為主）
    erp_product_name = str(
        get_cell(group_rows_[0], "erp_product_name", "") or ""
    ).strip()
    seller_name = erp_product_name or group_name.strip()

    items, item_missing = build_items(
        group_rows_=group_rows_,
        field_map=field_map,
        item_defaults=item_defaults,
        validation=validation,
        img_host=img_host,
        fallback_main_image_url=fallback_main_image_url,
        selected_image_folder=selected_image_folder,
    )

    payload: Dict[str, Any] = {
        payload_field: int(category_code),
        "sellerProductName": seller_name,
        "vendorId": cfg.vendor_id,
        "items": items,
    }

    payload.update(payload_defaults)
    normalize_shipping_fee_tw(payload)
    normalize_tw_payload(payload)

    payload.setdefault("displayProductName", seller_name)
    payload.setdefault("generalProductName", seller_name)
    if brand:
        payload.setdefault("brand", brand)

    return payload, item_missing


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--dry-run", action="store_true", help="Only write payloads, do not call API"
    )
    ap.add_argument(
        "--fallback-main-image",
        default="",
        help="Fallback image URL if Excel main image is empty",
    )
    ap.add_argument(
        "--img-base-url",
        default="",
        help="TryCloudflare base url, e.g. https://xxxx.trycloudflare.com",
    )
    ap.add_argument("--assets-dir", default="assets", help="Local assets directory")
    ap.add_argument(
        "--dump-input", action="store_true", help="Dump parsed input rows then exit"
    )
    ap.add_argument(
        "--image-folder", default="", help="Subfolder under assets to use for images"
    )
    ap.add_argument(
        "--delivery-charge",
        type=int,
        default=None,
        metavar="NT",
        help="Override 運費 (長+寬+高級距)，如 65,70,90,105,135,180,285,365",
    )
    ap.add_argument(
        "--return-charge",
        type=int,
        default=None,
        metavar="NT",
        help="退貨運費 (NT)，預設可為運費 100%%",
    )
    ap.add_argument(
        "--delivery-charge-type",
        default="",
        choices=["NOT_FREE", "CONDITIONAL_FREE", "FREE"],
        help="運費模式：NOT_FREE / CONDITIONAL_FREE / FREE；空字串代表沿用 mapping.json",
    )
    ap.add_argument(
        "--free-over",
        type=int,
        default=None,
        metavar="NT",
        help="滿額免運門檻 (freeShipOverAmount)，僅在 delivery-charge-type=CONDITIONAL_FREE 時使用",
    )
    ap.add_argument(
        "--gsheet-header-row",
        type=int,
        default=None,
        metavar="N",
        help="覆寫 mapping 的 gsheet.header_row（預設常為 2）",
    )
    ap.add_argument(
        "--gsheet-data-start-row",
        type=int,
        default=None,
        metavar="N",
        help="覆寫 gsheet.data_start_row（資料從第幾列開始，如 4）",
    )
    ap.add_argument(
        "--gsheet-data-row-count",
        type=int,
        default=None,
        metavar="N",
        help="覆寫讀取筆數；指定時優先於 mapping 的 data_row_count，且不再套用 data_end_row",
    )
    args = ap.parse_args()

    if not CONFIG_PATH.exists():
        raise FileNotFoundError(f"Missing {CONFIG_PATH.resolve()}")

    cfg = load_config()
    mp = load_mapping()

    # -----------------------------
    # mapping.json -> variables
    # -----------------------------
    market = str(mp.get("market", "TW")).strip().upper()

    cat = mp.get("category", {}) or {}
    categories_json_path = str(cat.get("categories_json_path", "")).strip()
    fallback_code = int(cat.get("fallback_code", 0) or 0)
    payload_field = (
        str(cat.get("payload_field", "displayCategoryCode")).strip()
        or "displayCategoryCode"
    )

    auto = cat.get("auto", {}) or {}
    auto_source = (
        str(auto.get("source", "group_name")).strip() or "group_name"
    )  # "group_name" / "itemName"

    fixed = cat.get("fixed", {}) or {}
    fixed_code = int(fixed.get("code", 0) or 0)
    fixed_path = str(fixed.get("path", "")).strip()
    category_mode = str(cat.get("mode", "keyword") or "keyword").strip().lower()
    llm_min_confidence = float(cat.get("llm_min_confidence", 0.55) or 0.55)
    llm_top_k = int(cat.get("llm_top_k", 10) or 10)

    # 若你未來想放到 mapping.json 控制 match/strategy，也先給預設值
    search_match = str(cat.get("match", "AND")).strip().upper()  # "AND" / "OR"
    search_strategy = (
        str(cat.get("strategy", "best")).strip().lower()
    )  # "best" / "deepest" / "first"

    if category_mode == "llm":
        if not categories_json_path:
            raise ValueError(
                'mapping.json 的 category.categories_json_path 為必填（mode="llm"）'
            )
        preload_index(Path(categories_json_path))

    grouping = mp.get("grouping", {}) or {}
    group_key_col = str(grouping.get("group_key_column", "")).strip()
    fallback_key_col = str(grouping.get("fallback_group_key_column", "")).strip()

    field_map = mp.get("field_map", {}) or {}
    payload_defaults = mp.get("payload_defaults", {}) or {}
    item_defaults = mp.get("item_defaults", {}) or {}
    validation = mp.get("validation", {}) or {}
    image_folder_column = str(
        mp.get("image_folder_column", "sale_product_name") or "sale_product_name"
    ).strip()

    # fallback main image from CLI
    fallback_main_image_url = str(args.fallback_main_image or "").strip()

    # optional: image hosting (local assets -> vendorPath URL)
    # 若沒有在 CLI 帶入 img-base-url，則嘗試從 start_image_hosting.py 寫入的 image_hosting.json 讀取
    img_host = None
    image_host_cfg_path = Path("image_hosting.json")
    img_base_url = str(args.img_base_url or "").strip()
    assets_dir_arg = str(args.assets_dir or "").strip()

    if image_host_cfg_path.exists():
        try:
            _cfg = json.loads(image_host_cfg_path.read_text(encoding="utf-8"))
        except Exception:
            _cfg = {}
        if not img_base_url:
            cfg_url = str(_cfg.get("img_base_url", "") or "").strip()
            if cfg_url:
                img_base_url = cfg_url
        if (not assets_dir_arg or assets_dir_arg == "assets") and _cfg.get(
            "assets_dir"
        ):
            assets_dir_arg = str(_cfg.get("assets_dir") or "").strip() or assets_dir_arg
        ap = str(_cfg.get("assets_path", "") or "").strip()
        if ap and (not assets_dir_arg or assets_dir_arg == "assets"):
            assets_dir_arg = ap

    # 將最終決定的 assets_dir 回寫到 args，讓後續邏輯一致
    args.assets_dir = assets_dir_arg or args.assets_dir

    _ad = Path(str(args.assets_dir or "").strip() or "assets")
    if not _ad.is_absolute():
        wr = Path(__file__).resolve().parent.parent
        sa = wr / "shared" / "assets"
        if sa.is_dir() and _ad.parts == ("assets",):
            args.assets_dir = str(sa.resolve())
        else:
            args.assets_dir = str(
                (Path(__file__).resolve().parent / _ad).resolve()
            )
    else:
        args.assets_dir = str(_ad.resolve())

    if img_base_url:
        img_host = LocalImageHosting(
            assets_dir=Path(args.assets_dir), base_url=img_base_url
        )

    selected_image_folder = str(args.image_folder or "").strip()
    # -----------------------------
    # Read input rows (xlsx or gsheet)
    # -----------------------------
    inp = mp.get("input", {}) or {}
    fmt = str(inp.get("format", "xlsx")).strip().lower()

    data_rows: List[Dict[str, Any]] = []

    if fmt == "gsheet":
        gs = _merge_gsheet_config(inp.get("gsheet", {}) or {})
        header_row = int(gs.get("header_row") or 2)
        data_start_row = int(gs.get("data_start_row") or 4)
        if getattr(args, "gsheet_header_row", None) is not None:
            header_row = int(args.gsheet_header_row)
        if getattr(args, "gsheet_data_start_row", None) is not None:
            data_start_row = int(args.gsheet_data_start_row)

        data_rows, _header_map = read_sheet_records(
            credentials_path=str(gs.get("cred_path", "")).strip(),
            spreadsheet_id=str(gs.get("spreadsheet_id", "")).strip(),
            worksheet_name=str(gs.get("worksheet_name", "")).strip(),
            header_row=header_row,
            data_start_row=data_start_row,
        )
        limit = gs.get("data_row_count", None)
        end_row = gs.get("data_end_row", None)

        if getattr(args, "gsheet_data_row_count", None) is not None:
            limit = int(args.gsheet_data_row_count)
            end_row = None

        if any(
            getattr(args, k, None) is not None
            for k in (
                "gsheet_header_row",
                "gsheet_data_start_row",
                "gsheet_data_row_count",
            )
        ):
            print(
                "[INFO] gsheet 列範圍（CLI 覆寫後）："
                f" header_row={header_row}, data_start_row={data_start_row},"
                f" data_row_count={limit}, data_end_row={end_row}"
            )

        if limit is not None:
            data_rows = data_rows[: int(limit)]

        if end_row is not None:
            # end_row 是工作表的實際列號，data_start_row=4 代表 data_rows[0] 對應 sheet 第4列
            end_row = int(end_row)
            start = int(data_start_row)
            keep = max(0, end_row - start + 1)
            data_rows = data_rows[:keep]
    else:
        xlsx_path = Path(str(inp.get("path", "test.xlsx")))
        skip_rows = [int(x) for x in (inp.get("skip_rows") or [])]
        header_row = int(inp.get("header_row") or 1)

        _, data_rows = read_xlsx_rows(
            xlsx_path=xlsx_path,
            skip_rows=skip_rows,
            header_row=header_row,
        )

        # -----------------------------
    # Debug: dump input rows then exit
    # -----------------------------
    if args.dump_input:
        print("============================================================")
        print("[DUMP] Parsed input rows")
        print(f"- format: {fmt}")
        print(f"- total rows: {len(data_rows)}")
        print("============================================================")

        # show first 10 rows, but only important columns if present
        preview_cols = [
            "sale_product_name",
            "erp_product_name",
            "erp_sku",
            "barcode",
            "sale_price",
            "list_price",
            "spec_name_1",
            "spec_value_1",
        ]

        for i, r in enumerate(data_rows[:10], start=1):
            slim = {k: r.get(k, "") for k in preview_cols if k in r}
            # fallback: if none of preview cols exist, dump whole row
            if not slim:
                slim = r
            print(f"[Row {i}] {json.dumps(slim, ensure_ascii=False)}")

        # also print suspicious rows (likely not products)
        bad = []
        for r in data_rows:
            sku = str(r.get("erp_sku", "") or "").strip()
            bc = str(r.get("barcode", "") or "").strip()
            name = str(r.get("erp_product_name", "") or "").strip()
            sale = str(r.get("sale_price", "") or "").strip()
            if (not name) or (not (sku or bc)) or (not sale):
                bad.append(r)

        print("------------------------------------------------------------")
        print(f"[DUMP] Suspicious rows (missing name/sku-or-barcode/price): {len(bad)}")
        for i, r in enumerate(bad[:10], start=1):
            print(
                f"[Bad {i}] "
                + json.dumps(
                    {
                        "sale_product_name": r.get("sale_product_name", ""),
                        "erp_product_name": r.get("erp_product_name", ""),
                        "erp_sku": r.get("erp_sku", ""),
                        "barcode": r.get("barcode", ""),
                        "sale_price": r.get("sale_price", ""),
                    },
                    ensure_ascii=False,
                )
            )

        print("done")
        return

    groups = group_rows(data_rows, group_key_col, fallback_key_col)

    report_rows: List[Dict[str, Any]] = []
    skipped_cat_report = _category_report_fields(resolved_mode="skipped", llm_final=None)

    for idx, (group_name, g_rows) in enumerate(groups.items(), start=1):
        group_image_folder = selected_image_folder
        if not group_image_folder:
            group_image_folder = str(
                get_cell(g_rows[0], image_folder_column, "") or ""
            ).strip()

        if img_host:
            if not group_image_folder:
                report_rows.append(
                    {
                        "group_name": group_name,
                        "category_code": "",
                        **skipped_cat_report,
                        "payload_file": "",
                        "uploaded": "NO",
                        "http_status": "",
                        "message": "skipped (missing image folder)",
                        "missing": (
                            "請在命令列指定 --image-folder，或在試算表欄位 "
                            f"'{image_folder_column}' 填入對應的 assets 子資料夾名稱"
                        ),
                    }
                )
                continue
            _asset_sub = Path(args.assets_dir) / group_image_folder
            if not _asset_sub.is_dir():
                report_rows.append(
                    {
                        "group_name": group_name,
                        "category_code": "",
                        **skipped_cat_report,
                        "payload_file": "",
                        "uploaded": "NO",
                        "http_status": "",
                        "message": "skipped (image folder not on disk)",
                        "missing": f"找不到資料夾：{_asset_sub}",
                    }
                )
                continue

        # -----------------------------
        # Per-group category resolve
        # -----------------------------
        query_text = group_name
        if auto_source == "itemName":
            item_name_col = str((field_map or {}).get("itemName", "")).strip()
            if item_name_col:
                v = get_cell(g_rows[0], item_name_col, "")
                if v is not None and str(v).strip():
                    query_text = str(v).strip()

        llm_audit_final: Optional[Dict[str, Any]] = None
        resolved_mode = ""

        if fixed_code > 0:
            category_code = fixed_code
            dbg = {"mode": "fixed", "fixed_code": fixed_code, "fixed_path": fixed_path}
            resolved_mode = "fixed"
        elif category_mode == "llm":
            llm_result = predict_category(
                build_product_for_llm(g_rows[0]),
                category_json_path=Path(categories_json_path),
                top_k=llm_top_k,
                min_confidence=llm_min_confidence,
            )
            fr_any = llm_result.get("final_result")
            fr = fr_any if isinstance(fr_any, dict) else {}
            llm_audit_final = fr
            code_parsed = _parse_display_category_code(fr.get("category_code"))
            llm_conf = float(fr.get("confidence") or 0)
            llm_ok = code_parsed is not None and llm_conf >= llm_min_confidence

            print("[CATEGORY][LLM] picked_code =", str(fr.get("category_code", "") or "").strip())
            print("[CATEGORY][LLM] picked_path =", str(fr.get("category_path", "") or "").strip())
            print("[CATEGORY][LLM] needs_review =", bool(fr.get("needs_review", False)))
            print("[CATEGORY][LLM] parse_mode =", fr.get("parse_mode", ""))
            print("[CATEGORY][LLM] reason =", fr.get("reason", ""))
            print("[CATEGORY][LLM] confidence =", fr.get("confidence", ""))
            raw_output = str(fr.get("raw_output", "") or "").strip()
            if raw_output:
                print("[CATEGORY][LLM] raw_output:")
                print(raw_output)

            top_k_candidates = llm_result.get("top_k_candidates") or []
            if isinstance(top_k_candidates, list) and top_k_candidates:
                print("[CATEGORY][LLM] top_k_candidates:")
                for j, c in enumerate(top_k_candidates[:10], start=1):
                    if isinstance(c, dict):
                        print(
                            "  "
                            f"{j}. code={c.get('code', '')} "
                            f"score={c.get('retrieval_score', 0)} "
                            f"path={c.get('path', '')}"
                        )

            if llm_ok:
                category_code = int(code_parsed)
                dbg = {
                    "mode": "llm",
                    "llm_result": fr,
                    "top_k_candidates": llm_result.get("top_k_candidates", []),
                    "query": query_text,
                }
                resolved_mode = "llm"
            else:
                category_code, dbg_kw = resolve_category_code(
                    categories_json_path=Path(categories_json_path),
                    fallback_code=fallback_code,
                    query=query_text,
                    match=search_match,
                    strategy=search_strategy,
                )
                dbg = {
                    **dbg_kw,
                    "mode": "llm_then_keyword",
                    "llm_attempt": fr,
                    "top_k_candidates": llm_result.get("top_k_candidates", []),
                    "query": query_text,
                }
                resolved_mode = "llm_then_keyword"
        else:
            category_code, dbg = resolve_category_code(
                categories_json_path=Path(categories_json_path),
                fallback_code=fallback_code,
                query=query_text,
                match=search_match,
                strategy=search_strategy,
            )
            resolved_mode = "keyword"

        cat_report = _category_report_fields(
            resolved_mode=resolved_mode,
            llm_final=llm_audit_final,
        )

        # write per-group debug candidates (optional)
        cand_path = OUT_DIR / f"category_candidates_{idx}.txt"
        with cand_path.open("w", encoding="utf-8") as f:
            f.write(f"mode={dbg.get('mode')}\n")
            f.write(f"query={dbg.get('query', query_text)}\n")
            f.write(f"fallback={fallback_code}\n")
            f.write(f"chosen={category_code}\n")
            if dbg.get("mode") == "fixed":
                f.write(f"fixed_path={dbg.get('fixed_path', '')}\n")
            if dbg.get("mode") in ("llm", "llm_then_keyword", "llm_fallback"):
                lr = dbg.get("llm_result") or dbg.get("llm_attempt")
                if isinstance(lr, dict):
                    f.write(
                        f"llm_category_code={lr.get('category_code', '')}\n"
                        f"llm_confidence={lr.get('confidence', '')}\n"
                        f"llm_reason={lr.get('reason', '')}\n"
                        f"llm_parse_mode={lr.get('parse_mode', '')}\n"
                    )
                    raw_output = str(lr.get("raw_output", "") or "").strip()
                    if raw_output:
                        f.write("\n--- llm_raw_output ---\n")
                        f.write(raw_output)
                        f.write("\n")
            f.write("\n")

            cands = dbg.get("candidates") or []
            for row in cands[:50]:
                if isinstance(row, (list, tuple)) and len(row) == 3:
                    f.write(f"{row[0]}\t{row[2]:.4f}\t{row[1]}\n")
                elif isinstance(row, (list, tuple)) and len(row) == 2:
                    f.write(f"{row[0]}\t{row[1]}\n")

            tk = dbg.get("top_k_candidates")
            if isinstance(tk, list) and tk:
                f.write("\n--- top_k_candidates (LLM retrieval) ---\n")
                for c in tk[:20]:
                    if isinstance(c, dict):
                        f.write(
                            f"{c.get('code', '')}\t{c.get('retrieval_score', 0)}\t"
                            f"{c.get('path', '')}\n"
                        )

        # -----------------------------
        # Build payload
        # -----------------------------
        payload, item_missing = build_payload(
            cfg=cfg,
            group_name=group_name,
            group_rows_=g_rows,
            category_code=category_code,
            payload_field=payload_field,
            field_map=field_map,
            payload_defaults=payload_defaults,
            item_defaults=item_defaults,
            validation=validation,
            img_host=img_host,
            fallback_main_image_url=fallback_main_image_url,
            selected_image_folder=group_image_folder,
        )

        # 運費：優先 CLI，否則依 gsheet 長+寬+高自動計算
        if getattr(args, "delivery_charge", None) is not None:
            apply_delivery_charge_override(payload, args.delivery_charge)
        else:
            dim_sum = get_dimension_sum(g_rows[0], field_map)
            if dim_sum is not None:
                apply_delivery_charge_override(payload, fee_from_dimension_sum(dim_sum))

        # 運費模式：NOT_FREE / CONDITIONAL_FREE / FREE（若未指定則沿用 mapping.json）
        if getattr(args, "delivery_charge_type", None):
            apply_shipping_mode_override(
                payload,
                args.delivery_charge_type,
                getattr(args, "free_over", None),
            )

        if getattr(args, "return_charge", None) is not None:
            apply_return_charge_override(payload, args.return_charge)

        payload_file = OUT_DIR / f"payload_{idx}.json"
        payload_file.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )

        missing: List[str] = []
        if not payload.get(payload_field):
            missing.append(payload_field)

        missing.extend(ensure_tw_validation(payload, validation))
        if item_missing:
            missing.append(
                f"images: {item_missing[:3]}{'...' if len(item_missing) > 3 else ''}"
            )

        if missing:
            report_rows.append(
                {
                    "group_name": group_name,
                    "category_code": payload.get(payload_field, ""),
                    **cat_report,
                    "payload_file": str(payload_file),
                    "uploaded": "NO",
                    "http_status": "",
                    "message": "skipped (missing required fields)",
                    "missing": "; ".join(missing),
                }
            )
            continue

        if args.dry_run:
            report_rows.append(
                {
                    "group_name": group_name,
                    "category_code": payload.get(payload_field, ""),
                    **cat_report,
                    "payload_file": str(payload_file),
                    "uploaded": "NO",
                    "http_status": "",
                    "message": "dry-run (payload only)",
                    "missing": "",
                }
            )
            continue

        status, text, resp_headers, req_headers = http_post(
            cfg, market=market, path=API_PATH, body=payload, timeout=60.0
        )

        resp_file = OUT_DIR / f"response_{idx}.json"
        resp_file.write_text(text, encoding="utf-8")

        (OUT_DIR / f"response_{idx}.headers.json").write_text(
            json.dumps(resp_headers, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (OUT_DIR / f"request_{idx}.headers.json").write_text(
            json.dumps(req_headers, ensure_ascii=False, indent=2), encoding="utf-8"
        )

        print(f"[HTTP] status={status}")
        print(f"[HTTP] body={text}")

        report_rows.append(
            {
                "group_name": group_name,
                "category_code": payload.get(payload_field, ""),
                **cat_report,
                "payload_file": str(payload_file),
                "uploaded": "YES" if 200 <= status < 300 else "NO",
                "http_status": status,
                "message": f"response saved: {resp_file}",
                "missing": "",
            }
        )

    write_report(report_rows)
    print(f"[OK] wrote: {REPORT_CSV}")
    print("done")


if __name__ == "__main__":
    main()
